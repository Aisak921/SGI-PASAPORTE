"""Datos base: 13 ubicaciones, admins iniciales y carga de inventario desde CSV.

Uso:  python -m app.seed data/inventario.csv
"""
import asyncio
import csv
import sys

# pyrefly: ignore [missing-import]
from sqlmodel import select

from .catalogo import TipoEquipo, Ubicacion, UBICACIONES_MONODISPOSITIVO, a_enum, norm_etiqueta
from .config import settings
from .database import SessionLocal, init_db
from .models import Activo, EstadoActivo, Estacion, Rol, Usuario


async def seed_base() -> None:
    """Inserta ubicaciones/admins que falten. Si vienes de una BD ya sembrada con nombres viejos
    (ej. 'Oficina del Jefe' antes de renombrarse a 'Jefe Pasaporte'), esto NO los renombra:
    corre `UPDATE estaciones_trabajo SET nombre='Jefe Pasaporte' WHERE nombre='Oficina del Jefe';` a mano una vez.
    """
    async with SessionLocal() as s:
        existentes = {e.nombre for e in (await s.exec(select(Estacion))).all()}
        for u in Ubicacion:
            if u.value not in existentes:
                s.add(Estacion(nombre=u.value, es_multiequipo=u not in UBICACIONES_MONODISPOSITIVO))
        for chat_id in settings.admin_ids:
            if not (await s.exec(select(Usuario).where(Usuario.chat_id == chat_id))).first():
                s.add(Usuario(chat_id=chat_id, nombre=f"Admin {chat_id}", rol=Rol.ADMIN.value))
        await s.commit()


async def cargar_inventario(ruta_archivo: str) -> int:
    """Carga activos desde un archivo CSV o Excel (.xlsx). Devuelve la cantidad de nuevos activos agregados."""
    async with SessionLocal() as s:
        estaciones = {e.nombre: e.id for e in (await s.exec(select(Estacion))).all()}
        nuevos = 0
        filas = []

        ruta_lower = str(ruta_archivo).lower()
        if ruta_lower.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(ruta_archivo, data_only=True)
            ws = wb["Inventario Detallado"] if "Inventario Detallado" in wb.sheetnames else wb.active
            headers = [str(c.value or "").strip().lower() for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                d = {headers[i]: (str(val).strip() if val is not None else "") for i, val in enumerate(row) if i < len(headers)}
                ub_str = d.get("ubicacion") or d.get("ubicación") or d.get("modulo") or d.get("módulo") or ""
                tipo_str = d.get("tipo_equipo") or d.get("tipo") or d.get("equipo") or ""
                etiqueta_str = d.get("etiqueta") or d.get("codigo") or d.get("código") or ""
                serial_str = d.get("serial") or d.get("serie") or d.get("s/n") or ""
                marca_str = d.get("marca_modelo") or d.get("marca") or d.get("modelo") or ""
                filas.append({
                    "ubicacion": ub_str,
                    "tipo_equipo": tipo_str,
                    "etiqueta": etiqueta_str,
                    "serial": serial_str,
                    "marca_modelo": marca_str,
                })
        else:
            with open(ruta_archivo, newline="", encoding="utf-8-sig") as f:
                for row in csv.DictReader(f):
                    filas.append(row)

        for fila in filas:
            ub = a_enum(Ubicacion, fila.get("ubicacion"))
            tipo = a_enum(TipoEquipo, fila.get("tipo_equipo"))
            if not ub or not tipo:
                print(f"Fila ignorada (ubicación o tipo inválido): {fila}")
                continue
            etiqueta = norm_etiqueta(fila.get("etiqueta") or fila.get("codigo"))
            if etiqueta and (await s.exec(select(Activo).where(Activo.etiqueta == etiqueta))).first():
                print(f"Etiqueta ya existe, se omite: {etiqueta}")
                continue
            serial = (fila.get("serial") or "").strip() or None
            if serial and (await s.exec(select(Activo).where(Activo.serial == serial))).first():
                print(f"Serial ya existe, se omite: {serial}")
                continue
            estado = EstadoActivo.EN_RESERVA if ub == Ubicacion.RECAMBIO else EstadoActivo.OPERATIVO
            s.add(
                Activo(
                    estacion_id=estaciones[ub.value],
                    tipo_equipo=tipo.value,
                    etiqueta=etiqueta,
                    serial=serial,
                    marca_modelo=(fila.get("marca_modelo") or "").strip() or None,
                    estado=estado.value,
                )
            )
            nuevos += 1
        await s.commit()
    print(f"Activos cargados: {nuevos}")
    return nuevos


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Uso: python -m app.seed <ruta_inventario.csv>")

    async def _run():
        await init_db()
        await seed_base()
        await cargar_inventario(sys.argv[1])

    asyncio.run(_run())
